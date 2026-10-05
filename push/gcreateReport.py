# gcreateReport.py — create a formatted .docx playtest report from DevBoard session data.
#
# Arg: {sheet_id}|{base64_json}
# JSON keys: game, client, my_name, my_company, my_address, my_email, my_phone, ref_code
# Returns: { ok, b64, filename, title }  (b64 = base64-encoded .docx bytes)

import sys, os, json, base64, socket, re, io, tempfile
from datetime import datetime

socket.setdefaulttimeout(20)

CRED_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'credentials.json')

# ── Parse args ───────────────────────────────────────────────────────────────
raw = sys.argv[1] if len(sys.argv) > 1 else ''
parts = raw.split('|', 1)
if len(parts) < 2:
    print(json.dumps({"error": "Expected: sheet_id|base64_json"}))
    sys.exit(1)

sheet_id = parts[0].strip()
try:
    data = json.loads(base64.b64decode(parts[1]).decode('utf-8'))
except Exception as e:
    print(json.dumps({"error": f"Bad payload: {e}"}))
    sys.exit(1)

game       = data.get('game',       '').strip()
client     = data.get('client',     '').strip()
my_name    = data.get('my_name',    '').strip()
my_company = data.get('my_company', '').strip()
my_address = data.get('my_address', '').strip()
my_email   = data.get('my_email',   '').strip()
my_phone   = data.get('my_phone',   '').strip()
ref_code   = data.get('ref_code',   '').strip()

if not game:
    print(json.dumps({"error": "game is required"}))
    sys.exit(1)

# ── Load session data (cache first, fall back to live sheet read) ─────────────
import gspread

base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
dev_file = os.path.join(base_dir, 'sheets', sheet_id, f'[{game}] dev.json')

if os.path.exists(dev_file):
    with open(dev_file, 'r', encoding='utf-8') as f:
        rows = json.load(f)
elif not os.path.isdir(os.path.join(base_dir, 'sheets', sheet_id)):
    print(json.dumps({"error": f"Sheet directory not found for id: {sheet_id}"}))
    sys.exit(1)
else:
    # No cached JSON — read directly from the Google Sheet
    if not os.path.exists(CRED_FILE):
        print(json.dumps({"error": "credentials.json not found"}))
        sys.exit(1)
    tab_name = f'[{game}] dev'
    try:
        sa = gspread.service_account(filename=CRED_FILE)
        wb = sa.open_by_key(sheet_id)
        ws = next((w for w in wb.worksheets() if w.title == tab_name), None)
        if ws is None:
            all_titles = [w.title for w in wb.worksheets()]
            print(json.dumps({"error": f"No dev tab found for: {game} (tabs: {all_titles})"}))
            sys.exit(1)
        all_vals = ws.get_all_values(value_render_option='FORMATTED_VALUE')
        if not all_vals:
            rows = []
        else:
            headers = all_vals[0]
            rows = [dict(zip(headers, r)) for r in all_vals[1:]]
    except Exception as e:
        print(json.dumps({"error": f"Sheet read failed: {e}"}))
        sys.exit(1)

# ── Build sessions (mirrors buildSessions() in devboard-common.js) ───────────
sessions = []
current  = None

for row in rows:
    date   = str(row.get('Date',         '') or '').strip()
    event  = str(row.get('Event',        '') or '').strip()
    people = str(row.get('People',       '') or '').strip()
    obs    = str(row.get('Observations', '') or row.get('Observation', '') or '').strip()
    sol    = str(row.get('Solutions',    '') or row.get('Thoughts', '') or row.get('Solution', '') or '').strip()

    if date or event:
        if event in ('Time', 'Material'):
            current = None
            continue
        label   = event + (' ' + people if people else '')
        current = {
            'date':     date,
            'testnum':  label,
            'event':    event,
            'location': obs,
            'length':   sol,
            'testers':  [],
            'obs':      [],
        }
        sessions.append(current)
    elif current:
        if people:
            if people.startswith('[sub:') and people.endswith(']'):
                current['submittedBy'] = people[5:-1]
            else:
                tname = people.split('@')[0].strip() if '@' in people else people
                current['testers'].append(tname)
        elif obs or sol:
            current['obs'].append({'obs': obs, 'sol': sol})

playtest_sessions = [s for s in sessions if s['obs'] or s.get('location')]

# ── Image helper ──────────────────────────────────────────────────────────────
IMAGE_RE = re.compile(r'=IMAGE\s*\(\s*["\']?(https?[^"\')\s]+)["\']?\s*\)', re.IGNORECASE)

def fetch_image(url):
    """Download image bytes; returns (bytes, ext) or (None, None) on failure."""
    try:
        import urllib.request
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as resp:
            data = resp.read()
            ct = resp.headers.get('Content-Type', '')
            ext = 'png' if 'png' in ct else 'jpg'
            return data, ext
    except Exception:
        return None, None

def split_obs_text(text):
    """Split text into segments: str or {'image_url': url}."""
    segments = []
    last = 0
    for m in IMAGE_RE.finditer(text):
        if m.start() > last:
            segments.append(text[last:m.start()])
        segments.append({'image_url': m.group(1)})
        last = m.end()
    if last < len(text):
        segments.append(text[last:])
    return segments

# ── Build .docx ───────────────────────────────────────────────────────────────
try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError as e:
    print(json.dumps({"error": f"python-docx not installed: {e}"}))
    sys.exit(1)

TEAL   = RGBColor(0x1a, 0x5f, 0x7a)
GRAY   = RGBColor(0x66, 0x66, 0x66)
BLACK  = RGBColor(0x11, 0x11, 0x11)

def add_run(para, text, bold=False, italic=False, size=None, color=None):
    run = para.add_run(text)
    run.bold   = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return run

def set_para_spacing(para, before=0, after=0, line=None):
    pf = para.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after  = Pt(after)
    if line:
        pf.line_spacing = Pt(line)

def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    set_para_spacing(p, before=4, after=4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1a5f7a')
    pBdr.append(bottom)
    pPr.append(pBdr)

today = datetime.today()
date_str = today.strftime('%b ') + str(today.day) + today.strftime(', %Y')

doc_title = f"{game} — Playtest Report"
if ref_code:
    doc_title += f" ({ref_code})"

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin    = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

# ── Header block ─────────────────────────────────────────────────────────────
header_text = my_company or my_name
if header_text:
    p = doc.add_paragraph()
    set_para_spacing(p, before=0, after=2)
    add_run(p, header_text, bold=True, size=20, color=TEAL)

# Date line
p = doc.add_paragraph()
set_para_spacing(p, before=0, after=1)
add_run(p, 'Date: ', bold=True, size=10, color=BLACK)
add_run(p, date_str, size=10, color=BLACK)

# Address lines
if my_address:
    for addr_line in my_address.replace('\r\n', '\n').split('\n'):
        addr_line = addr_line.strip()
        if addr_line:
            p = doc.add_paragraph()
            set_para_spacing(p, before=0, after=0)
            add_run(p, addr_line, size=10, color=GRAY)

# Prepared by
if my_name and my_company:
    p = doc.add_paragraph()
    set_para_spacing(p, before=2, after=0)
    add_run(p, 'Prepared by: ', bold=True, size=10, color=BLACK)
    add_run(p, my_name, size=10, color=BLACK)

add_horizontal_rule(doc)

# ── CLIENT / GAME block ───────────────────────────────────────────────────────
tbl = doc.add_table(rows=1, cols=2)
tbl.style = 'Table Grid'
tbl.columns[0].width = Cm(5)
tbl.columns[1].width = Cm(10)

c0, c1 = tbl.rows[0].cells
p0 = c0.paragraphs[0]
add_run(p0, 'CLIENT\n', bold=True, size=9, color=TEAL)
add_run(p0, client or '—', size=10, color=BLACK)

p1 = c1.paragraphs[0]
add_run(p1, 'GAME\n', bold=True, size=9, color=TEAL)
add_run(p1, game, size=10, color=BLACK)

doc.add_paragraph()  # spacer

# ── TESTS section ─────────────────────────────────────────────────────────────
p = doc.add_paragraph()
set_para_spacing(p, before=6, after=4)
add_run(p, 'TESTS', bold=True, size=14, color=TEAL)

if playtest_sessions:
    for s in playtest_sessions:
        # Session heading
        heading = s['testnum']
        if s['date']:
            heading += f"  —  {s['date']}"
        p = doc.add_paragraph()
        set_para_spacing(p, before=8, after=2)
        add_run(p, heading, bold=True, size=12, color=BLACK)

        # Testers
        if s['testers']:
            p = doc.add_paragraph()
            set_para_spacing(p, before=0, after=1)
            add_run(p, 'Testers: ', bold=True, size=10, color=GRAY)
            add_run(p, ', '.join(s['testers']), italic=True, size=10, color=GRAY)

        # Location / Length meta
        meta_parts = []
        if s.get('location'):
            meta_parts.append(s['location'])
        if s.get('length'):
            meta_parts.append(s['length'])
        if meta_parts:
            p = doc.add_paragraph()
            set_para_spacing(p, before=0, after=3)
            add_run(p, ' · '.join(meta_parts), italic=True, size=10, color=GRAY)

        # Observations + solutions
        for entry in s['obs']:
            obs_text = (entry.get('obs') or '').strip()
            sol_text = (entry.get('sol') or '').strip()

            if not obs_text and not sol_text:
                continue

            # Split obs_text into text / image segments
            segments = split_obs_text(obs_text) if obs_text else []

            # Collect image segments separately; build plain text for bullet
            plain_parts = []
            image_urls  = []
            for seg in segments:
                if isinstance(seg, dict):
                    image_urls.append(seg['image_url'])
                else:
                    plain_parts.append(seg.strip())

            plain_text = ' '.join(p for p in plain_parts if p)

            # Bullet paragraph for the observation text
            if plain_text:
                p = doc.add_paragraph(style='List Bullet')
                set_para_spacing(p, before=1, after=1)
                add_run(p, plain_text, size=10, color=BLACK)

            # Solution as indented sub-bullet
            if sol_text:
                # Split sol too for any images
                sol_segs    = split_obs_text(sol_text)
                sol_plain   = ' '.join(seg.strip() for seg in sol_segs if isinstance(seg, str) and seg.strip())
                sol_imgs    = [seg['image_url'] for seg in sol_segs if isinstance(seg, dict)]
                if sol_plain:
                    p = doc.add_paragraph(style='List Bullet 2')
                    set_para_spacing(p, before=0, after=1)
                    add_run(p, sol_plain, size=10, color=GRAY)
                for img_url in sol_imgs:
                    img_bytes, _ = fetch_image(img_url)
                    if img_bytes:
                        p = doc.add_paragraph()
                        set_para_spacing(p, before=2, after=2)
                        p.paragraph_format.left_indent = Cm(1.5)
                        run = p.add_run()
                        run.add_picture(io.BytesIO(img_bytes), width=Inches(4))

            # Observation images
            for img_url in image_urls:
                img_bytes, _ = fetch_image(img_url)
                if img_bytes:
                    p = doc.add_paragraph()
                    set_para_spacing(p, before=2, after=4)
                    run = p.add_run()
                    run.add_picture(io.BytesIO(img_bytes), width=Inches(5))

else:
    p = doc.add_paragraph()
    set_para_spacing(p, before=4, after=4)
    add_run(p, 'No playtest sessions recorded yet.', italic=True, size=10, color=GRAY)

# ── Footer ────────────────────────────────────────────────────────────────────
add_horizontal_rule(doc)
footer_parts = []
if my_company:
    footer_parts.append(my_company)
elif my_name:
    footer_parts.append(my_name)
if my_phone:
    footer_parts.append(my_phone)
if my_email:
    footer_parts.append(my_email)
if footer_parts:
    p = doc.add_paragraph()
    set_para_spacing(p, before=2, after=0)
    add_run(p, '  ·  '.join(footer_parts), size=9, color=GRAY)

# ── Serialize to base64 ───────────────────────────────────────────────────────
buf = io.BytesIO()
doc.save(buf)
b64 = base64.b64encode(buf.getvalue()).decode('ascii')

safe_game = re.sub(r'[^\w\s-]', '', game).strip().replace(' ', '_')
if ref_code:
    filename = f"{ref_code}_{safe_game}_Playtest_Report.docx"
else:
    filename = f"{safe_game}_Playtest_Report.docx"

print(json.dumps({"ok": True, "b64": b64, "filename": filename, "title": doc_title}))
sys.exit(0)
